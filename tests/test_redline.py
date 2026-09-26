import asyncio
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from fastapi import HTTPException
from contract_web.store import Store
from contract_web.redline import Redline, validate_docx


class RedlineTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.env=patch.dict(os.environ,{'CW_REDLINE_ENABLED':'1'});self.env.start()
        self.store=Store(self.root);uid=self.store.add_user('alice','test-password-long')
        self.store.execute("UPDATE users SET preferences=json_set(preferences,'$.redline_enabled',json('true')) WHERE id=?",(uid,))
        self.u=self.store.one('SELECT * FROM users WHERE id=?',(uid,))
        self.store.execute('INSERT INTO workspaces(id,user_id,document_id,title,created) VALUES(?,?,?,?,?)',('workspace',uid,'doc','合同',1))
        self.t={'id':'thread','workspace_id':'workspace'}
        self.store.execute('INSERT INTO documents(id,user_id,workspace_id,thread_id,filename,suffix,source_hash) VALUES(?,?,?,?,?,?,?)',('doc',uid,'workspace',None,'合同.docx','.docx','original-hash'))
        self.d=self.store.one('SELECT * FROM documents WHERE id=?',('doc',))
        folder=self.store.user_root(uid)/'sources/doc';folder.mkdir(parents=True)
        doc=Document();doc.add_heading('采购合同',0);doc.add_paragraph('付款期限为30天。');doc.add_paragraph('其他条款保持不变。')
        doc.save(folder/'source.docx');self.original=(folder/'source.docx').read_bytes()
        self.s=Redline(self.store);await self.s.ensure(self.d,self.u)

    async def asyncTearDown(self):
        self.env.stop();self.tmp.cleanup()

    def body(self,action,**kw):
        return {'document_id':'doc','base_version':self.s.version(self.d)['id'],'request_id':os.urandom(12).hex(),'action':action,**kw}

    async def change(self):
        return await self.s.apply(self.u,self.t,self.body('apply',agent=True,changes=[{'quote':'30天','replacement':'60天','reason':'用户要求'}]))

    async def test_roundtrip_decisions_and_original(self):
        old=self.s.version(self.d)['id'];receipt=await self.change()
        self.assertTrue(receipt['saved']);self.assertEqual(receipt['parent_id'],old)
        state=self.s.state(self.d);changes=state['snapshot']['changes'];self.assertTrue(changes)
        self.assertIn('60天',state['snapshot']['text'])
        await self.s.apply(self.u,self.t,self.body('decide',ids=[c['id'] for c in changes],decision='reject'))
        self.assertIn('30天',self.s.state(self.d)['snapshot']['text'])
        self.assertFalse(self.s.state(self.d)['snapshot']['changes'])
        self.assertEqual(self.original,(self.s.directory(self.d).parent/'source.docx').read_bytes())

    async def test_failed_batch_never_commits(self):
        before=self.s.version(self.d)['id']
        with self.assertRaises(ValueError):
            await self.s.apply(self.u,self.t,self.body('apply',changes=[{'quote':'30天','replacement':'60天'},{'quote':'不存在','replacement':'新内容'}]))
        self.assertEqual(before,self.s.version(self.d)['id'])

    async def test_idempotency_and_stale_version(self):
        body=self.body('apply',changes=[{'quote':'30天','replacement':'60天'}])
        a=await self.s.apply(self.u,self.t,body);b=await self.s.apply(self.u,self.t,body)
        self.assertEqual(a,b)
        with self.assertRaises(HTTPException) as cm:
            await self.s.apply(self.u,self.t,{**body,'request_id':'another-request'})
        self.assertEqual(cm.exception.status_code,409)
        self.assertFalse(self.store.one('SELECT yield_requested FROM redline_documents')['yield_requested'])

    async def test_clean_export_and_restore(self):
        original=self.s.version(self.d)['id'];await self.change()
        with self.assertRaises(ValueError): await self.s.apply(self.u,self.t,self.body('export',clean=True))
        await self.s.apply(self.u,self.t,self.body('restore',restore_version=original))
        self.assertNotEqual(self.s.version(self.d)['id'],original)
        receipt=await self.s.apply(self.u,self.t,self.body('export',clean=True))
        out=self.store.user_root(self.u['id'])/'published'/receipt['artifact_id']/'content.docx'
        validate_docx(out.read_bytes(),clean=True)
        self.assertTrue(receipt['saved'])

    async def test_comment_reply_and_clean_copy(self):
        await self.s.apply(self.u,self.t,self.body('comment',operation='create',quote='付款期限',text='请确认'))
        state=self.s.state(self.d);c=state['snapshot']['comments'][0]
        await self.s.apply(self.u,self.t,self.body('comment',operation='reply',id=c['id'],text='已确认'))
        await self.s.apply(self.u,self.t,self.body('comment',operation='resolve',id=c['id']))
        receipt=await self.s.apply(self.u,self.t,self.body('export',clean=True))
        self.assertTrue(self.s.state(self.d)['snapshot']['comments'])
        out=self.store.user_root(self.u['id'])/'published'/receipt['artifact_id']/'content.docx'
        validate_docx(out.read_bytes(),clean=True)

    async def test_browser_handoff_can_save_without_deadlock(self):
        lease={'client_id':'browser-session'};self.s.lease(self.d,self.t,lease)
        task=asyncio.create_task(self.change());await asyncio.sleep(.25)
        self.assertTrue(self.s.lease(self.d,self.t,lease)['yield_requested'])
        data=self.s.file(self.d).read_bytes()
        await asyncio.wait_for(self.s.apply(self.u,self.t,self.body('save',client_id=lease['client_id'],content_hash=hashlib.sha256(data).hexdigest()),data=data),10)
        self.s.lease(self.d,self.t,{**lease,'release':True})
        receipt=await task
        self.assertTrue(receipt['saved']) # An unchanged browser flush must not invalidate the Agent's base.

    async def test_attachment_isolation_and_tampered_file(self):
        self.store.execute('UPDATE documents SET thread_id=? WHERE id=?',('another','doc'))
        with self.assertRaises(HTTPException):self.s.authorize(self.u,self.t,'doc')
        self.store.execute('UPDATE documents SET thread_id=NULL WHERE id=?',('doc',))
        self.s.file(self.d).write_bytes(b'broken')
        with self.assertRaises(HTTPException):self.s.file(self.d)

    async def test_cancel_handoff_releases_flag(self):
        self.s.lease(self.d,self.t,{'client_id':'browser-session'})
        task=asyncio.create_task(self.change());await asyncio.sleep(.1);task.cancel()
        with self.assertRaises(asyncio.CancelledError):await task
        self.assertFalse(self.store.one('SELECT yield_requested FROM redline_documents')['yield_requested'])

    async def test_reused_id_mismatch_releases_handoff(self):
        body=self.body('apply',changes=[{'quote':'30天','replacement':'60天'}]);await self.s.apply(self.u,self.t,body)
        with self.assertRaises(HTTPException):await self.s.apply(self.u,self.t,{**body,'summary':'different'})
        self.assertFalse(self.store.one('SELECT yield_requested FROM redline_documents')['yield_requested'])

    async def test_discussion_survives_rejected_anchor_and_restart(self):
        await self.change()
        await self.s.apply(self.u,self.t,self.body('comment',operation='create',quote='6',text='请确认这个修改'))
        changes=self.s.state(self.d)['snapshot']['changes']
        await self.s.apply(self.u,self.t,self.body('decide',ids=[c['id'] for c in changes],decision='reject'))
        restarted=Redline(self.store);state=restarted.state(self.d)
        discussions=state['snapshot']['comments']+state['detached_comments']
        self.assertTrue(any(c['text']=='请确认这个修改' for c in discussions))
        if state['detached_comments']:
            c=state['detached_comments'][0]
            body=self.body('discussion',id=c['id'],resolved=True)
            a=await restarted.apply(self.u,self.t,body);b=await restarted.apply(self.u,self.t,body)
            self.assertEqual(a,b);self.assertEqual(restarted.state(self.d)['detached_comments'][0]['status'],'resolved')

    async def test_saved_versions_survive_restart_with_real_diff_and_group(self):
        before=self.s.version(self.d)['id'];receipt=await self.change();restarted=Redline(self.store)
        state=restarted.state(self.d);self.assertEqual(state['version']['id'],receipt['version_id'])
        self.assertTrue(all(c['authorEmail']==receipt['author_email'] for c in state['snapshot']['changes']))
        diff=await restarted.apply(self.u,self.t,{'action':'diff','document_id':'doc','before_version':before,'after_version':receipt['version_id']})
        self.assertIn('-付款期限为30天。',diff['diff']);self.assertIn('+付款期限为60天。',diff['diff'])

    async def test_insert_delete_and_table_cell(self):
        await self.s.apply(self.u,self.t,self.body('apply',changes=[{'type':'insert_before','quote':'付款期限','replacement':'买方的'},{'type':'insert_after','quote':'30天','replacement':'内支付'},{'type':'delete','quote':'其他条款保持不变。'}]))
        text=self.s.state(self.d)['snapshot']['text'];self.assertIn('买方的付款期限为30天内支付',text);self.assertNotIn('其他条款',text)
        doc=Document();table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='单价为100元。'
        doc.save(self.s.directory(self.d).parent/'source.docx')
        path=self.root/'table.docx';doc.save(path)
        out=self.root/'table-edited.docx'
        result=await self.s.engine(path,'apply',output=str(out),author={'name':'Agent','email':'agent@example.test'},changes=[{'quote':'100元','replacement':'200元'}])
        self.assertIn('200元',result['text']);self.assertTrue(result['changes'])

    async def test_receipt_contains_real_diff_and_immutable_download(self):
        before=self.s.version(self.d)['id'];receipt=await self.change()
        self.assertIn('-付款期限为30天。',receipt['diff'])
        self.assertIn('+付款期限为60天。',receipt['diff'])
        self.assertIn(receipt['version_id'],receipt['download_url'])
        self.assertIn('thread_id=thread',receipt['download_url'])
        self.assertGreater(receipt['pending_changes'],0)
        self.assertEqual(receipt['parent_id'],before)

    async def test_export_saved_version_while_new_version_is_leased(self):
        before=self.s.version(self.d)['id'];await self.change()
        self.s.lease(self.d,self.t,{'client_id':'another-browser'})
        body=self.body('export',base_version=before)
        receipt=await asyncio.wait_for(self.s.apply(self.u,self.t,body),3)
        self.assertEqual(receipt,await self.s.apply(self.u,self.t,body))
        self.assertEqual(receipt['version_id'],before)
        out=self.store.user_root(self.u['id'])/'published'/receipt['artifact_id']/'content.docx'
        self.assertEqual(out.read_bytes(),self.s.file(self.d,before).read_bytes())
        self.assertNotEqual(self.s.version(self.d)['id'],before)
        self.assertFalse(self.store.one('SELECT yield_requested FROM redline_documents')['yield_requested'])

    async def test_repacked_identical_docx_does_not_create_version_but_format_edit_does(self):
        self.s.lease(self.d,self.t,{'client_id':'browser-session'})
        old=self.s.version(self.d)['id'];buf=io.BytesIO()
        with ZipFile(io.BytesIO(self.original)) as source,ZipFile(buf,'w') as output:
            for name in source.namelist():output.writestr(name,source.read(name))
        data=buf.getvalue()
        body=self.body('save',client_id='browser-session',content_hash=hashlib.sha256(data).hexdigest())
        receipt=await self.s.apply(self.u,self.t,body,data=data)
        self.assertTrue(receipt['unchanged']);self.assertEqual(receipt['version_id'],old)
        self.assertEqual(receipt,await self.s.apply(self.u,self.t,body,data=data))
        doc=Document(io.BytesIO(data));doc.paragraphs[1].runs[0].bold=True
        buf=io.BytesIO();doc.save(buf);data=buf.getvalue()
        with patch.object(self.s,'engine',wraps=self.s.engine) as inspect:
            changed=await self.s.apply(self.u,self.t,self.body('save',client_id='browser-session',content_hash=hashlib.sha256(data).hexdigest()),data=data)
            self.assertEqual(inspect.await_count,1) # Uploaded DOCX needs one authoritative read.
        self.assertNotEqual(changed['version_id'],old)
        self.assertEqual(changed['diff'],'') # Formatting must survive despite an empty text diff.

    async def test_unsupported_paragraph_merge_is_atomic_and_explicit(self):
        old=self.s.version(self.d)['id']
        with self.assertRaisesRegex(ValueError,'UNSUPPORTED_PARAGRAPH_MERGE'):
            await self.s.apply(self.u,self.t,self.body('apply',changes=[
                {'quote':'30天','replacement':'60天'},
                {'quote':'付款期限为30天。\n其他条款保持不变。','replacement':'合并后的内容'}]))
        self.assertEqual(old,self.s.version(self.d)['id'])
        self.assertFalse(self.s.state(self.d)['snapshot']['changes'])

    async def test_pending_deletions_not_searchable_and_missing_search_does_not_leak_raw_text(self):
        await self.change()
        for quote,count in [('付款期限为60天。',1),('30天',0),('不存在',0)]:
            result=await self.s.apply(self.u,self.t,{'action':'find','document_id':'doc','quote':quote})
            self.assertEqual(result['matches']['total'],count)
            self.assertEqual(result['text_view'],'visible')

        await self.s.apply(self.u,self.t,self.body('apply',changes=[{'quote':'付款期限为60天。','replacement':'付款期限为90天。'}]))
        state=self.s.state(self.d)
        self.assertIn('付款期限为90天。',state['snapshot']['text'])
        self.assertIn('其他条款保持不变。',state['snapshot']['text'])
        info=await self.s.apply(self.u,self.t,{'action':'inspect','document_id':'doc'})
        self.assertFalse(info['capabilities']['tracked_paragraph_merge'])
        self.assertNotIn('versions',info);self.assertNotIn('operations',info)
        self.assertIn('versions',await self.s.apply(self.u,self.t,{'action':'inspect','document_id':'doc','history':True}))

    async def test_visible_fallback_keeps_duplicate_detection(self):
        await self.s.apply(self.u,self.t,self.body('apply',changes=[{'quote':'其他条款保持不变。','replacement':'付款期限为60天。'}]))
        await self.change()
        found=await self.s.apply(self.u,self.t,{'action':'find','document_id':'doc','quote':'付款期限为60天。'})
        self.assertEqual(found['matches']['total'],2)
        old=self.s.version(self.d)['id']
        with self.assertRaisesRegex(ValueError,'AMBIGUOUS_TARGET'):
            await self.s.apply(self.u,self.t,self.body('apply',changes=[{'quote':'付款期限为60天。','replacement':'付款期限为90天。'}]))
        self.assertEqual(old,self.s.version(self.d)['id'])

    async def test_agent_request_identity_is_automatic_and_replayable(self):
        body=self.body('apply',agent=True,changes=[{'quote':'30天','replacement':'60天'}]);body.pop('request_id')
        a=await self.s.apply(self.u,self.t,body);b=await self.s.apply(self.u,self.t,body)
        self.assertEqual(a,b);self.assertTrue(a['operation_id'].startswith('agent-'))
        self.assertNotIn('request_id',body)
        with self.assertRaises(HTTPException):
            await self.s.apply(self.u,self.t,{**body,'changes':[{'quote':'30天','replacement':'90天'}]})

    async def test_version_activity_and_label_are_separate_from_document_versions(self):
        from contract_web.redline import version_activity
        old=self.s.version(self.d);before=self.s.file(self.d).read_bytes()
        result=self.s.label_version(self.d,old['id'],{'label':'  内部审核完成稿  ','revision':0})
        self.assertEqual(result['label'],'内部审核完成稿')
        self.assertEqual(self.s.label_version(self.d,old['id'],{'label':'内部审核完成稿','revision':0}),result)
        self.assertEqual(self.s.version(self.d)['id'],old['id']);self.assertEqual(self.s.file(self.d).read_bytes(),before)
        self.assertEqual(len(self.s.state(self.d)['versions']),1)
        with self.assertRaises(HTTPException) as conflict:self.s.label_version(self.d,old['id'],{'label':'过期的名称','revision':0})
        self.assertEqual(conflict.exception.status_code,409)
        self.s.label_version(self.d,old['id'],{'label':'','revision':1})
        self.assertEqual(self.s.state(self.d)['versions'][0]['label'],'')
        self.assertEqual(version_activity({'action':'save','review_events':[{'decision':'accept','ids':['a','b']},{'decision':'reject','ids':['c']}]},{},{}),('review','接受 2 处修订，拒绝 1 处修订',{'accept':2,'reject':1}))
        await self.change()
        v=self.s.state(self.d)['versions'][0];self.assertEqual(v['kind'],'agent');self.assertEqual(v['thread_id'],self.t['id']);self.assertEqual(v['actor_id'],self.u['id'])

    async def test_legacy_history_migration_preserves_versions_and_receipts(self):
        await self.change()
        old=self.store.all('SELECT * FROM redline_versions ORDER BY created')
        receipts=self.store.all('SELECT * FROM redline_operations ORDER BY created')
        self.store.execute('DROP TABLE redline_version_meta')
        migrated=Redline(self.store)
        self.assertEqual(old,self.store.all('SELECT * FROM redline_versions ORDER BY created'))
        self.assertEqual(receipts,self.store.all('SELECT * FROM redline_operations ORDER BY created'))
        history=migrated.state(self.d)['versions']
        self.assertEqual(history[0]['kind'],'agent');self.assertEqual(history[-1]['kind'],'initial')
        migrated.label_version(self.d,history[0]['id'],{'label':'保留标记','revision':0})
        self.assertEqual(Redline(self.store).state(self.d)['versions'][0]['label'],'保留标记')


if __name__=='__main__':unittest.main()
