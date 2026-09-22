import io,json,os,unittest,hashlib,importlib.util,tempfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from tests import test_workbench as workbench_tests

class RedlineApiTests(unittest.TestCase):
    def setUp(self):
        self.env=patch.dict(os.environ,{'CW_REDLINE_USERS':'alice'});self.env.start()
        self.base=workbench_tests.WorkbenchTests();self.base.setUp();self.c=self.base.client
        d=Document();d.add_paragraph('付款期限为30天。');b=io.BytesIO();d.save(b)
        self.w=self.c.post('/api/workspaces',content=b.getvalue(),headers={**self.base.headers,'X-Filename':'test.docx'}).json()
        self.t=self.c.post('/api/workspaces/'+self.w['id']+'/threads',json={},headers=self.base.headers).json()
        self.did=self.w['document_id'];self.url=f'/api/redline/{self.did}?thread_id={self.t["id"]}'
    def tearDown(self):self.base.tearDown();self.env.stop()
    def test_native_transport_and_binary_download(self):
        state=self.c.get(self.url);self.assertEqual(state.status_code,200)
        token=(self.base.store.user_root(self.base.uid)/'threads'/self.t['id']/'.publish-token').read_text()
        body={'kind':'redline','action':'apply','document_id':self.did,'base_version':state.json()['version']['id'],'request_id':'native-operation-1','changes':[{'quote':'30天','replacement':'60天'}]}
        r=self.c.post('/internal/artifacts',json=body,headers={'Authorization':'Bearer '+token});self.assertEqual(r.status_code,200,r.text);self.assertTrue(r.json()['saved'])
        direct=self.c.get(r.json()['download_url']);self.assertEqual(direct.status_code,200)
        self.assertEqual(hashlib.sha256(direct.content).hexdigest(),r.json()['hash'])
        direct_url=r.json()['download_url']
        body.update(action='export',request_id='native-export-1',base_version=r.json()['version_id']);body.pop('changes')
        exported=self.c.post('/internal/artifacts',json=body,headers={'Authorization':'Bearer '+token}).json()
        downloaded=self.c.get(exported['download_url']);self.assertEqual(downloaded.status_code,200)
        self.assertEqual(hashlib.sha256(downloaded.content).hexdigest(),exported['hash'])
        self.assertIn('wordprocessingml',downloaded.headers['content-type'])
        meta=self.c.get('/api/artifacts/'+exported['artifact_id']).json();self.assertTrue(meta['redline']);self.assertEqual(meta['formats'],['docx'])
        self.base.login('bob');self.assertEqual(self.c.get(direct_url).status_code,404);self.assertEqual(self.c.get(self.url).status_code,404);self.assertEqual(self.c.get(exported['download_url']).status_code,404)
    def test_flag_and_invalid_write_do_not_expose_editing(self):
        with patch.dict(os.environ,{'CW_REDLINE_USERS':''}):self.assertEqual(self.c.get(self.url).status_code,403)
        bad=self.c.post(self.url.replace('?', '/operations?'),json={'action':'save'},headers=self.base.headers)
        self.assertEqual(bad.status_code,422)
    def test_foreign_thread_attachment_stays_inaccessible(self):
        d=Document();d.add_paragraph('附件');b=io.BytesIO();d.save(b)
        attached=self.c.post('/api/threads/'+self.t['id']+'/attachments',content=b.getvalue(),headers={**self.base.headers,'X-Filename':'attachment.docx'}).json()
        other=self.c.post('/api/workspaces/'+self.w['id']+'/threads',json={},headers=self.base.headers).json()
        self.assertEqual(self.c.get(f'/api/redline/{attached["id"]}?thread_id={other["id"]}').status_code,404)
        own=self.c.get(f'/api/redline/{attached["id"]}?thread_id={self.t["id"]}').json()
        label_url=f'/api/redline/{attached["id"]}/versions/{own["version"]["id"]}/label?thread_id={other["id"]}'
        self.assertEqual(self.c.patch(label_url,json={'label':'不可串用','revision':0},headers=self.base.headers).status_code,404)
    def test_version_label_scope_and_immutable_file(self):
        state=self.c.get(self.url).json();vid=state['version']['id']
        path=f'/api/redline/{self.did}/versions/{vid}/label?thread_id={self.t["id"]}'
        file_url=f'/api/redline/{self.did}/file?thread_id={self.t["id"]}&version_id={vid}'
        before=self.c.get(file_url).content
        r=self.c.patch(path,json={'label':'客户沟通前版本','revision':0},headers=self.base.headers)
        self.assertEqual(r.status_code,200,r.text)
        self.assertEqual(self.c.get(file_url).content,before)
        state=self.c.get(self.url).json();self.assertEqual(len(state['versions']),1);self.assertEqual(state['versions'][0]['label'],'客户沟通前版本')
        self.assertEqual(self.c.patch(path,json={'label':'另一名称','revision':0},headers=self.base.headers).status_code,409)
        self.assertEqual(self.c.patch(path.replace(vid,'not-this-document'),json={'label':'错误版本','revision':0},headers=self.base.headers).status_code,404)
        self.assertEqual(self.c.patch(path,json={'label':'x'*81,'revision':1},headers=self.base.headers).status_code,422)
        self.base.login('bob');self.assertEqual(self.c.patch(path,json={'label':'越权修改','revision':1},headers=self.base.headers).status_code,404)
