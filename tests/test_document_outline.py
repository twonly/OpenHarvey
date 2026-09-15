import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from contract_web.documents import prepare
from contract_web.document_outline import document_outline
from contract_web.pdf_preview import read_pdf_outline


def outlined_pdf(action=False):
    """A real two-page PDF with nested bookmarks and text, without extra libraries."""
    stream1 = b'BT /F1 12 Tf 40 740 Td (Payment terms. Payment is due within thirty days after acceptance.) Tj 0 -20 Td (The supplier provides services and technical support during the term.) Tj ET'
    stream2 = b'BT /F1 12 Tf 40 740 Td (Delivery and acceptance. The Payment schedule follows the contract.) Tj ET'
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R /Outlines 8 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 5 0 R >> >> /Contents 6 0 R >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 5 0 R >> >> /Contents 7 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length ' + str(len(stream1)).encode() + b' >>\nstream\n' + stream1 + b'\nendstream',
        b'<< /Length ' + str(len(stream2)).encode() + b' >>\nstream\n' + stream2 + b'\nendstream',
        b'<< /Type /Outlines /First 9 0 R /Last 9 0 R /Count 2 >>',
        b'<< /Title (Payment terms) /Parent 8 0 R /Dest [3 0 R /Fit] /First 10 0 R /Last 10 0 R /Count 1 >>',
        b'<< /Title (Delivery and acceptance) /Parent 9 0 R /Dest [4 0 R /Fit] >>',
    ]
    if action:
        objects[9] = objects[9].replace(b'/Dest [4 0 R /Fit]', b'/A << /S /GoTo /D [4 0 R /Fit] >>')
    return pdf_objects(objects)


def pdf_objects(objects):
    data = b'%PDF-1.4\n';offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data));data += f'{number} 0 obj\n'.encode() + obj + b'\nendobj\n'
    xref = len(data)
    data += f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode()
    data += b''.join(f'{offset:010} 00000 n \n'.encode() for offset in offsets[1:])
    return data + f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode()


def linked_pdf(toc=True):
    stream = (b'BT /F1 12 Tf 40 740 Td (' + (b'Contents' if toc else b'Related references') +
              b') Tj 0 -30 Td (1 Payment terms .... 1) Tj 0 -25 Td (1.1 Delivery .... 2) Tj ET')
    body = b'BT /F1 12 Tf 40 740 Td (Delivery and acceptance terms.) Tj ET'
    return pdf_objects([
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R 4 0 R] /Count 2 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 5 0 R >> >> /Contents 6 0 R /Annots [8 0 R 9 0 R 10 0 R 11 0 R 12 0 R] >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 5 0 R >> >> /Contents 7 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream',
        b'<< /Length '+str(len(body)).encode()+b' >>\nstream\n'+body+b'\nendstream',
        b'<< /Type /Annot /Subtype /Link /Rect [40 706 210 722] /Dest [4 0 R /Fit] >>',
        # Duplicate annotation for the same row.
        b'<< /Type /Annot /Subtype /Link /Rect [40 706 210 722] /A << /S /GoTo /D [4 0 R /Fit] >> >>',
        # Only the final page number is clickable: recover its row title.
        b'<< /Type /Annot /Subtype /Link /Rect [126 681 137 697] /A << /S /GoTo /D [4 0 R /Fit] >> >>',
        b'<< /Type /Annot /Subtype /Link /Rect [40 736 210 752] /A << /S /URI /URI (https://example.com) >> >>',
        b'<< /Type /Annot /Subtype /Link /Rect [40 736 210 752] /A << /S /GoToR /F (external.pdf) /D [0 /Fit] >> >>',
    ])


class DocumentOutlineTests(unittest.TestCase):
    def test_printed_toc_links_deduplicate_and_use_actual_destinations(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'source.pdf';path.write_bytes(linked_pdf())
            self.assertEqual(read_pdf_outline(path), [])
            self.assertEqual(prepare(path, 'aaaaaaaaaaaa')['outline'], {
                'source': 'pdf-toc-links', 'entries': [
                    {'title': '1 Payment terms', 'level': 1, 'page': 2},
                    {'title': '1.1 Delivery', 'level': 2, 'page': 2}]})
            path.write_bytes(linked_pdf(toc=False))
            self.assertEqual(prepare(path, 'aaaaaaaaaaaa')['outline']['entries'], [])

    def test_word_native_inherited_headings_body_override_and_legacy_maps(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'source.docx';doc = Document()
            doc.add_paragraph('第一条 普通编号正文，不是标题')
            doc.add_heading('付款安排', 1)
            doc.add_table(rows=2, cols=2).cell(0, 0).text = '表格内容'
            style = doc.styles.add_style('Company Heading', WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = doc.styles['Heading 2']
            doc.add_paragraph('验收条件', style)
            p = doc.add_paragraph('不应出现在目录的正文', 'Heading 1')
            value = OxmlElement('w:outlineLvl');value.set(qn('w:val'), '9');p._p.get_or_add_pPr().append(value)
            doc.save(path);mapped = prepare(path, 'aaaaaaaaaaaa')
            entries = mapped['outline']['entries']
            self.assertEqual([(e['title'], e['level'], e['block_id']) for e in entries], [('付款安排', 1, 'B1'), ('验收条件', 2, 'B5')])
            before = (Path(root)/'contract.md').read_bytes()
            for segment in mapped['segments']:
                segment.pop('outline_level', None)
            mapped.pop('outline')
            self.assertEqual(document_outline(path, mapped)['entries'], entries)
            self.assertEqual((Path(root)/'contract.md').read_bytes(), before)

    def test_pdf_bookmarks_keep_hierarchy_and_real_page_targets(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root)/'source.pdf';path.write_bytes(outlined_pdf())
            mapped = prepare(path, 'aaaaaaaaaaaa')
            self.assertEqual(mapped['outline'], {'source': 'pdf-bookmarks', 'entries': [
                {'title': 'Payment terms', 'level': 1, 'page': 1},
                {'title': 'Delivery and acceptance', 'level': 2, 'page': 2}]})
            self.assertEqual(read_pdf_outline(Path(__file__).parent/'fixtures/text-contract.pdf'), [])
            path.write_bytes(outlined_pdf(action=True))
            self.assertEqual(read_pdf_outline(path), mapped['outline']['entries'])

    def test_markdown_headings_ignore_code_and_txt_does_not_infer(self):
        with tempfile.TemporaryDirectory() as root:
            text = '# 主合同\n正文付款。\n```md\n## 代码中的标题\n```\n## 交付 ###\n第一条 不是标题\n'
            path = Path(root)/'source.md';path.write_text(text)
            entries = prepare(path, 'aaaaaaaaaaaa')['outline']['entries']
            self.assertEqual([(e['title'], e['block_id']) for e in entries], [('主合同', 'B0'), ('交付', 'B5')])
            path = Path(root)/'source.txt';path.write_text(text)
            self.assertEqual(prepare(path, 'aaaaaaaaaaaa')['outline']['entries'], [])
