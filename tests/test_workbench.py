import time
import copy
import io
import json
import secrets
import tempfile
import threading
import unittest
from unittest.mock import patch
from pathlib import Path

from fastapi.testclient import TestClient
from docx import Document

from contract_web.app import create_app
from contract_web.documents import prepare, validate_citations, read_blocks
from contract_web.runtime import Runtime, visible_event, visible_messages
from contract_web.store import Store


class FakeRuntime(Runtime):
    """Transport fake: the Web must still enforce ownership and save real files."""
    servers = {}

    def __init__(self, config):
        super().__init__(config)
        self.server = self.servers.setdefault(config["url"], {"sessions": {}, "calls": [], "requests": {"question": [], "permission": []}})
        self.server.setdefault("skills", [{"name": name, "description": "业务说明", "location": f'{config["skill_root"]}/skills/{name}/SKILL.md', "content": "# 业务 Skill\n正文"}
                                          for name in ("contract-summary", "contract-review")])

    async def call(self, method, path, *, tid=None, body=None, params=None, allow_not_found=False):
        self.server["calls"].append((method, path, tid, body))
        if path == "/config":
            return {"model": "glm/glm-test", "provider": {"glm": {"options": {"apiKey": "PRIVATE_KEY"}, "models": {"glm-test": {}}},
                     "deepseek": {"models": {"flash": {}}}}}
        if path == "/provider":
            return {"connected": ["glm", "deepseek"], "all": [
                {"id": "glm", "models": {"glm-test": {"name": "GLM"}, "unconfigured": {"name": "Hidden"}}},
                {"id": "deepseek", "models": {"flash": {"name": "DeepSeek"}}}]}
        if path == "/skill":
            return self.server["skills"]
        if path == "/instance/dispose":
            return True
        if path == "/path":
            return {"worktree": "/", "directory": self.directory(tid)}
        if path == "/session" and method == "POST":
            sid = "ses_"+secrets.token_hex(12)
            self.server["sessions"][sid] = {"id": sid, "title": "新对话", "messages": [], "status": {"type": "idle"}, "permission": body["permission"]}
            return {"id": sid}
        if path == "/session/status":
            return {sid: s["status"] for sid, s in self.server["sessions"].items()}
        if path in {"/question", "/permission"}:
            return self.server["requests"][path[1:]]
        if path.startswith(("/question/", "/permission/")):
            return True
        parts = path.strip("/").split("/")
        if len(parts) == 2 and method == "DELETE":
            if self.server.pop("delete_fail", False):
                raise RuntimeError("native delete failed")
            self.server["sessions"].pop(parts[1], None)
            return True
        session = self.server["sessions"][parts[1]]
        if len(parts) == 2:
            if method == "PATCH":
                session.update(body)
            return session
        if parts[2] == "message":
            messages = session["messages"]
            if params and params.get('before'):
                messages = messages[:next(i for i,m in enumerate(messages) if m['info']['id']==params['before'])]
            return messages[-params['limit']:] if params and params.get('limit') else messages
        if parts[2] == "todo":
            return []
        if parts[2] == "abort":
            session["status"] = {"type": "idle"}
            return True
        if parts[2] == "prompt_async":
            session["status"] = {"type": "busy"}
            session["messages"].append({"info": {"id": body.get("messageID") or "msg_"+secrets.token_hex(6), "role": "user"}, "parts": body["parts"]})
            return None
        raise AssertionError((method, path))

    async def events(self, thread):
        yield {"type": "session.status", "properties": {"sessionID": "other", "status": {"type": "busy"}}}
        yield {"type": "session.status", "properties": {"sessionID": thread["session_id"], "status": {"type": "idle"}}}
        yield {"type": "message.part.updated", "properties": {"sessionID": thread["session_id"], "part": {"type": "reasoning", "text": "PRIVATE_REASONING"}}}


class WorkbenchTests(unittest.TestCase):
    def setUp(self):
        FakeRuntime.servers = {}
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.library = self.root/"library"
        self.library.mkdir()
        (self.library/"rules.yaml").write_text("- id: TEST-1\n  industry: 通用\n  name: 付款\n  baseline: 按期付款\n")
        self.app = create_app(self.root, FakeRuntime, self.library)
        self.store = self.app.state.store
        self.uid = self.store.add_user("alice", "alice-password-2026")
        self.bid = self.store.add_user("bob", "bob-password-2026")
        legacy=self.store.one("SELECT id FROM config_items WHERE kind='risk' AND org_id='default'")
        self.store.execute('UPDATE config_items SET owner_id=?,updated=? WHERE id=?',(self.uid,time.time()+1,legacy['id']))
        self.store.execute('UPDATE users SET risk_scheme=? WHERE id=?',(legacy['id'],self.uid))
        configs = {name: {"url": "http://"+name, "password": "runtime-password", "work_root": "/work", "source_root": "/sources", "skill_root": "/opt/contract"} for name in ("alice", "bob")}
        (self.root/"runtimes.json").write_text(json.dumps(configs))
        self.client = TestClient(self.app)
        self.headers = {"X-Workbench-Request": "1"}
        self.login("alice")

    def tearDown(self):
        self.client.close()
        self.app.state.outlines.close()
        self.tmp.cleanup()

    def login(self, name):
        response = self.client.post("/api/login", json={"username": name, "password": name+"-password-2026"}, headers=self.headers)
        self.assertEqual(response.status_code, 200, response.text)

    def make_workspace(self, text="甲方：采购公司\n乙方：服务公司\n合同总金额 128000 元。\n验收后 30 日付款。"):
        w = self.client.post("/api/workspaces", content=text.encode(), headers={**self.headers, "X-Filename": "contract.txt"})
        self.assertEqual(w.status_code, 200, w.text)
        w = w.json()
        t = self.client.post(f'/api/workspaces/{w["id"]}/threads', json={}, headers=self.headers)
        self.assertEqual(t.status_code, 200, t.text)
        return w, t.json()

    def test_legacy_outline_is_read_without_rewriting_source_and_stays_scoped(self):
        w, t = self.make_workspace()
        doc = Document();doc.add_heading('付款条款', 1);doc.add_paragraph('验收后付款。')
        file = io.BytesIO();doc.save(file)
        attached = self.client.post(f'/api/threads/{t["id"]}/attachments', content=file.getvalue(),
                                    headers={**self.headers, 'X-Filename': 'terms.docx'}).json()
        docid = attached['id']
        outline_url = f'/api/documents/{docid}/outline?thread_id={t["id"]}'
        self.wait_outline(outline_url)
        path = self.store.user_root(self.uid)/'sources'/docid/'document.json'
        (path.parent/'outline.json').unlink()
        mapping = json.loads(path.read_text());mapping.pop('outline')
        for segment in mapping['segments']:
            segment.pop('outline_level', None)
        path.write_text(json.dumps(mapping));before = path.read_bytes()
        url = f'/api/documents/{docid}?thread_id={t["id"]}'
        response = self.client.get(url)
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(self.wait_outline(outline_url)['entries'][0]['block_id'], 'B0')
        self.assertEqual(response.json()['segments'], mapping['segments'])
        self.assertEqual(path.read_bytes(), before)
        other = self.client.post(f'/api/workspaces/{w["id"]}/threads', json={}, headers=self.headers).json()
        self.assertEqual(self.client.get(f'/api/documents/{docid}?thread_id={other["id"]}').status_code, 404)
        self.login('bob');self.assertEqual(self.client.get(url).status_code, 404)
        self.assertEqual(self.client.get(outline_url).status_code, 404)

    def wait_outline(self, url):
        for _ in range(100):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 200, response.text)
            outline = response.json()['outline']
            if outline.get('status') != 'pending':
                return outline
            time.sleep(.01)
        self.fail('Background outline did not finish')

    def test_outline_does_not_block_upload_or_read_and_is_deduplicated(self):
        from contract_web.document_outline import document_outline
        started, release = threading.Event(), threading.Event()
        def delayed(path, mapped):
            started.set()
            release.wait(5)
            return document_outline(path, mapped)
        with patch('contract_web.outline_jobs.document_outline', side_effect=delayed) as extract:
            try:
                response = self.client.post('/api/workspaces', content=b'# Payment\nNet 30',
                    headers={**self.headers, 'X-Filename': 'terms.md'})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertTrue(started.wait(1))
                docid = response.json()['document_id'];url = f'/api/documents/{docid}'
                first = self.client.get(url).json()
                self.assertEqual(first['outline']['status'], 'pending')
                self.assertTrue(first['segments'])
                for _ in range(3):
                    self.assertEqual(self.client.get(url+'/outline').json()['outline']['status'], 'pending')
                self.assertEqual(extract.call_count, 1)
            finally:
                release.set()
            result = self.wait_outline(url+'/outline')
            self.assertEqual(result['entries'][0]['title'], 'Payment')
            path = self.store.user_root(self.uid)/'sources'/docid
            self.assertEqual(json.loads((path/'document.json').read_text())['outline']['status'], 'pending')
            self.assertEqual(self.client.get(url).json()['outline'], result)
            self.assertEqual(extract.call_count, 1)
            self.login('bob')
            self.assertEqual(self.client.get(url+'/outline').status_code, 404)

    def send(self, t, text="合同价款是多少？", **kwargs):
        return self.client.post(f'/api/threads/{t["id"]}/messages', json={"text": text, **kwargs}, headers=self.headers)

    def finish(self,t,error=None):
        self.native(t)['messages'].append({'info':{'id':'msg_'+secrets.token_hex(6),'role':'assistant','finish':'stop','time':{'completed':1},**({'error':error} if error else {})},'parts':[{'type':'text','text':'完成'}]})
        self.native(t)['status']={'type':'idle'}

    def native(self, t):
        return FakeRuntime.servers["http://alice"]["sessions"][t["session_id"]]

    def read_all(self, w, t):
        path = self.store.user_root(self.uid)/"sources"/w["document_id"]/"contract.md"
        self.native(t)["messages"].append({"info": {"id": "msg_read", "role": "assistant"}, "parts": [
            {"id": "p1", "type": "tool", "tool": "read", "state": {"status": "completed", "input": {"filePath": f'/sources/{w["document_id"]}/contract.md'}, "output": path.read_text()}}]})

    def report(self, w):
        m = json.loads((self.store.user_root(self.uid)/"sources"/w["document_id"]/"document.json").read_text())
        return {"kind": "summary", "title": "合同摘要", "source_hash": m["source_hash"], "content": f'# 合同摘要\n\n金额为 128000 元。{m["segments"][2]["citation"]}'}, m

    def publish(self, t, body):
        token = (self.store.user_root(self.uid)/"threads"/t["id"]/".publish-token").read_text()
        return self.client.post("/internal/artifacts", json=body, headers={"Authorization": "Bearer "+token})

    def test_auth_required_csrf_and_disabled_account(self):
        self.client.cookies.clear()
        self.assertEqual(self.client.get("/api/workspaces").status_code, 401)
        self.assertEqual(self.client.post("/api/login", json={}).status_code, 403)
        self.assertEqual(self.client.post("/api/login", json={}, headers={**self.headers, "Origin": "https://evil.example"}).status_code, 403)
        self.login("alice")
        self.store.execute("UPDATE users SET active=0 WHERE id=?", (self.uid,))
        self.assertEqual(self.client.get("/api/workspaces").status_code, 401)

    def test_native_skill_catalog_discovers_new_business_skill_and_selects_it(self):
        w, t = self.make_workspace()
        server = FakeRuntime.servers["http://alice"]
        server["skills"] += [
            {"name": "contract-change-risk", "description": "变更审查", "location": "/opt/contract/skills/contract-change-risk/SKILL.md", "content": "# 合同变更风险\nPRIVATE_BODY"},
            {"name": "customize-opencode", "location": "/internal/skills/customize-opencode/SKILL.md", "content": "PRIVATE_BODY"},
            {"name": "outside", "location": "/opt/contract/skills/../outside/SKILL.md", "content": "PRIVATE_BODY"}]
        response = self.client.get(f'/api/threads/{t["id"]}')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual({s["name"] for s in response.json()["skills"]}, {"contract-change-risk", "contract-summary", "contract-review"})
        self.assertNotIn("PRIVATE_BODY", response.text)
        self.assertNotIn("/opt/contract", response.text)
        self.assertEqual(self.send(t, skill="contract-change-risk").status_code, 202)
        prompt = next(body for method, path, _, body in reversed(server["calls"]) if path.endswith("prompt_async"))
        self.assertIn("明确选择 contract-change-risk", prompt["system"])
        self.assertTrue(prompt["parts"][0]["text"].startswith("/contract-change-risk\n"))
        snapshot = self.client.get(f'/api/threads/{t["id"]}').json()
        self.assertEqual(snapshot["messages"][0]["parts"][0]["text"], "/contract-change-risk\n合同价款是多少？")

    def test_model_selection_persists_per_thread_and_reaches_native_prompt(self):
        w, t = self.make_workspace()
        catalog = self.client.get("/api/models")
        self.assertEqual(catalog.json()["selected"], "glm/glm-test")
        self.assertEqual(len(catalog.json()["models"]), 2)
        self.assertNotIn("PRIVATE_KEY", catalog.text)
        self.assertNotIn("unconfigured", catalog.text)
        route = f"/api/threads/{t['id']}/model"
        self.assertEqual(self.client.put(route, json={"model": "deepseek/flash"}, headers=self.headers).status_code, 200)
        self.login("bob")
        self.assertEqual(self.client.get("/api/models").json()["selected"], "glm/glm-test")
        self.login("alice")
        self.assertEqual(self.client.get("/api/models",params={"thread_id":t["id"]}).json()["selected"], "deepseek/flash")
        other = self.client.post(f"/api/workspaces/{w['id']}/threads",json={},headers=self.headers).json()
        self.assertEqual(self.client.get("/api/models",params={"thread_id":other["id"]}).json()["selected"], "glm/glm-test")
        self.assertEqual(self.send(t).status_code, 202)
        prompt = next(c[3] for c in reversed(FakeRuntime.servers["http://alice"]["calls"]) if c[1].endswith("prompt_async"))
        self.assertEqual(prompt["model"], {"providerID": "deepseek", "modelID": "flash"})
        self.native(t)["status"] = {"type": "idle"}
        self.assertEqual(self.send(t, model="glm/unconfigured").status_code, 422)
        self.assertEqual(self.client.put(route, json={"model": "missing"}, headers=self.headers).status_code, 422)
        self.assertEqual(self.send(t, model="glm/glm-test").status_code, 202)

    def test_unchanged_skills_do_not_dispose_native_instance(self):
        import yaml
        w, t = self.make_workspace()
        root = Path(__file__).resolve().parents[1] / "runtime/skills"
        skills = []
        for path in root.rglob("SKILL.md"):
            _, front, content = path.read_text().split("---", 2)
            meta = yaml.safe_load(front)
            skills.append({**meta, "content": content, "location": "/opt/contract/skills/" + str(path.relative_to(root))})
        server = FakeRuntime.servers["http://alice"]
        server["skills"] = skills
        route = f'/api/threads/{t["id"]}/skills/refresh'
        self.assertEqual(self.client.post(route, json={}, headers=self.headers).status_code, 200)
        self.assertFalse(any(c[1] == "/instance/dispose" for c in server["calls"]))
        server["skills"][0]["content"] = "outdated"
        self.assertEqual(self.client.post(route, json={}, headers=self.headers).status_code, 200)
        self.assertTrue(any(c[1] == "/instance/dispose" for c in server["calls"]))

    def test_citation_locations_describe_only_current_materials(self):
        w, t = self.make_workspace()
        snap = self.client.get(f'/api/threads/{t["id"]}').json()
        self.assertEqual(snap["documents"][0]["locations"]["B2"], {"page": None, "ordinal": 3})
        self.login("bob")
        self.assertEqual(self.client.get(f'/api/threads/{t["id"]}').status_code, 404)

    def test_skill_refresh_is_owned_idle_scoped_and_preserves_history(self):
        w, t = self.make_workspace()
        route = f'/api/threads/{t["id"]}/skills/refresh'
        self.assertEqual(self.send(t).status_code, 202)
        self.assertEqual(self.client.post(route, json={}, headers=self.headers).status_code, 409)
        self.native(t)["status"] = {"type": "idle"}
        before = copy.deepcopy(self.native(t)["messages"])
        self.assertEqual(self.client.post(route, json={}, headers=self.headers).status_code, 200)
        self.assertEqual(self.native(t)["messages"], before)
        self.assertIn(("POST", "/instance/dispose", t["id"], None), FakeRuntime.servers["http://alice"]["calls"])
        self.assertEqual(self.send(t, skill="missing-skill").status_code, 422)
        self.assertEqual(self.send(t, skill="customize-opencode").status_code, 422)
        self.login("bob")
        self.assertEqual(self.client.post(route, json={}, headers=self.headers).status_code, 404)

    def test_upload_and_thread_do_not_copy_conversations_to_web_db(self):
        w, t = self.make_workspace()
        self.assertNotIn("save_token", t)
        self.assertEqual(self.send(t).status_code, 202)
        state = self.client.get(f'/api/threads/{t["id"]}').json()
        self.assertEqual(state["messages"][0]["parts"][0]["text"], "合同价款是多少？")
        names = self.store.all("SELECT name FROM sqlite_master WHERE type='table'")
        self.assertNotIn("messages", [x["name"] for x in names])
        self.assertEqual(state["status"]["type"], "busy")
        self.assertEqual(self.send(t).status_code, 202)
        self.assertEqual(self.client.post(f'/api/threads/{t["id"]}/abort', headers=self.headers).status_code, 200)
        self.assertEqual(self.client.get(f'/api/threads/{t["id"]}').json()["status"]["type"], "idle")

    def test_all_intents_pass_to_native_runtime(self):
        _, t = self.make_workspace()
        for text in ["你好", "好的", "什么是人工智能", "告诉我金额，再生成摘要", "/contract-summary", "乙方是谁？"]:
            if self.native(t)["messages"]:self.finish(t)
            self.assertEqual(self.send(t, text).status_code, 202)
            self.assertEqual(self.native(t)["messages"][-1]["parts"][0]["text"], text)
        self.finish(t)
        self.assertEqual(self.send(t, "做一份摘要", skill="contract-summary").status_code, 202)
        calls = FakeRuntime.servers["http://alice"]["calls"]
        self.assertIn("contract-summary", calls[-1][3]["system"])

    def test_cross_user_every_resource_denied(self):
        w,t = self.make_workspace();self.read_all(w,t)
        body,_ = self.report(w);a = self.publish(t,body)
        self.assertEqual(a.status_code,200,a.text)
        aid=a.json()["artifact_id"]
        self.login("bob")
        self.assertEqual(self.client.get("/api/workspaces").json(),[])
        for path in [f'/api/workspaces/{w["id"]}', f'/api/threads/{t["id"]}', f'/api/threads/{t["id"]}/events',
                     f'/api/documents/{w["document_id"]}?thread_id={t["id"]}',
                     f'/api/documents/{w["document_id"]}/file?thread_id={t["id"]}',
                     f'/api/artifacts/{aid}', f'/api/artifacts/{aid}/file']:
            self.assertEqual(self.client.get(path).status_code,404,path)
        self.assertEqual(self.send(t).status_code,404)
        self.assertEqual(self.client.post(f'/api/threads/{t["id"]}/abort',headers=self.headers).status_code,404)

    def test_thread_attachment_not_visible_to_another_thread(self):
        w,t = self.make_workspace()
        t2 = self.client.post(f'/api/workspaces/{w["id"]}/threads',headers=self.headers,json={}).json()
        a = self.client.post(f'/api/threads/{t["id"]}/attachments', content=b'PRIVATE_ATTACHMENT',headers={**self.headers,"X-Filename":"private.txt"}).json()
        state = self.client.get(f'/api/threads/{t2["id"]}').json()
        self.assertEqual(len(state["documents"]),1)
        self.assertEqual(self.client.get(f'/api/documents/{a["id"]}?thread_id={t2["id"]}').status_code,404)
        self.assertFalse(any(a["id"] in r["pattern"] for r in self.native(t2)["permission"]))
        self.assertTrue(any(a["id"] in r["pattern"] for r in self.native(t)["permission"]))

    def test_real_docx_export_and_idempotent_persistence(self):
        w,t=self.make_workspace();self.read_all(w,t);body,_=self.report(w)
        saved=self.publish(t,body)
        self.assertEqual(saved.status_code,200,saved.text)
        aid=saved.json()["artifact_id"]
        self.assertEqual(self.publish(t,body).json()["artifact_id"],aid)
        response=self.client.get(f'/api/artifacts/{aid}/file?format=docx')
        self.assertEqual(response.status_code,200)
        doc=Document(io.BytesIO(response.content))
        self.assertIn("128000",'\n'.join(p.text for p in doc.paragraphs))
        new_app=create_app(self.root,FakeRuntime,self.library)
        with TestClient(new_app) as restarted:
            restarted.cookies.update(self.client.cookies)
            self.assertEqual(restarted.get(f'/api/artifacts/{aid}').json()["content"],body["content"])
            self.assertEqual(restarted.get(f'/api/threads/{t["id"]}').status_code,200)

    def test_native_read_coverage_invalid_citations_and_versions(self):
        w,t=self.make_workspace();body,m=self.report(w)
        self.assertEqual(self.publish(t,body).status_code,422)
        self.read_all(w,t)
        self.native(t)["messages"][0]["parts"][0]["state"]["output"]='truncated search result'
        self.assertEqual(self.publish(t,body).status_code,422)
        self.native(t)["messages"].clear();self.read_all(w,t)
        wrong={**body,"content":body["content"]+'【Dffffffffffff:B0】'}
        self.assertEqual(self.publish(t,wrong).status_code,422)
        self.assertEqual(self.publish(t,{**body,"source_hash":"stale"}).status_code,422)
        self.assertEqual(self.store.all("SELECT * FROM artifacts"),[])

    def test_review_complete_rule_list_and_source_quotes(self):
        w,t=self.make_workspace();self.read_all(w,t);body,m=self.report(w)
        body.update(kind="review",findings=[],content=body["content"]+" TEST-1")
        self.assertEqual(self.publish(t,body).status_code,422)
        body["findings"]=[{"risk_id":"TEST-1","verdict":"命中","reason":"付款期限存在风险","evidence":[{
            "document_id":w["document_id"],"source_hash":m["source_hash"],"block_id":"B3","quote":"不存在的条款"}]}]
        self.assertEqual(self.publish(t,body).status_code,422)
        body["findings"][0]["evidence"][0]["quote"]="验收后 30 日付款。"
        self.assertEqual(self.publish(t,body).status_code,200)

    def test_sse_scopes_session_and_omits_reasoning(self):
        _,t=self.make_workspace()
        response=self.client.get(f'/api/threads/{t["id"]}/events')
        self.assertIn(t["session_id"],response.text)
        self.assertNotIn('"other"',response.text)
        self.assertNotIn('PRIVATE_REASONING',response.text)

    def test_provider_response_headers_are_never_sent_to_browser(self):
        error={"name":"APIError","data":{"message":"quota exhausted","statusCode":402,
            "responseHeaders":{"set-cookie":"private-provider-cookie"},"responseBody":"private upstream body"}}
        message={"info":{"id":"m1","sessionID":"s1","error":error},"parts":[]}
        self.assertNotIn("private",json.dumps(visible_messages([message])))
        for event in [{"type":"session.error","properties":{"sessionID":"s1","error":error}},
                      {"type":"message.updated","properties":{"info":message["info"]}}]:
            clean=visible_event(event,"s1",set())
            self.assertNotIn("private",json.dumps(clean))
            self.assertIn("quota exhausted",json.dumps(clean))

    def test_native_compaction_events_and_summary_are_forwarded(self):
        updated={'type':'message.updated','properties':{'info':{'id':'summary-1','sessionID':'s1','role':'assistant','summary':True,'mode':'compaction'}}}
        self.assertEqual(visible_event(updated,'s1',set()),updated)
        part={'type':'message.part.updated','properties':{'part':{'id':'part-1','messageID':'summary-1','sessionID':'s1','type':'text','text':'原生压缩摘要'}}}
        self.assertEqual(visible_event(part,'s1',set()),part)
        compacted={'type':'session.compacted','properties':{'sessionID':'s1'}}
        self.assertEqual(visible_event(compacted,'s1',set()),compacted)
        self.assertEqual(visible_messages([{'info':updated['properties']['info'],'parts':[part['properties']['part']]}])[0]['parts'][0]['text'],'原生压缩摘要')

    def test_revision_requires_real_target_read_and_preserves_original(self):
        w,t=self.make_workspace();body,m=self.report(w)
        body={"kind":"revision","source_hash":m["source_hash"],"changes":[{
            "document_id":w["document_id"],"block_id":"B3","quote":"验收后 30 日付款。",
            "replacement":"验收后 15 日付款。","reason":"按用户要求缩短账期"}]}
        self.assertEqual(self.publish(t,body).status_code,422)
        self.read_all(w,t)
        self.native(t)['messages'][0]['parts'][0]['state']['output']=m['segments'][3]['citation']+' '+m['segments'][3]['text']
        response=self.publish(t,body);self.assertEqual(response.status_code,200,response.text)
        aid=response.json()['artifact_id'];artifact=self.client.get('/api/artifacts/'+aid).json()
        self.assertIn('15 日付款',artifact['content']);self.assertIn('30 日付款',artifact['content'])
        original=self.client.get(f'/api/documents/{w["document_id"]}/file?thread_id={t["id"]}')
        self.assertIn('30 日付款',original.text);self.assertNotIn('15 日付款',original.text)
        doc=Document(io.BytesIO(self.client.get(f'/api/artifacts/{aid}/file?format=docx').content))
        self.assertIn('15 日付款','\n'.join(p.text for p in doc.paragraphs))
        body['changes'][0]['quote']='虚构原条款';self.assertEqual(self.publish(t,body).status_code,422)

    def test_missing_published_file_is_reported_instead_of_download_success(self):
        w,t=self.make_workspace();self.read_all(w,t);body,_=self.report(w)
        aid=self.publish(t,body).json()['artifact_id']
        (self.store.user_root(self.uid)/'published'/aid/'report.docx').unlink()
        self.assertEqual(self.client.get(f'/api/artifacts/{aid}/file').status_code,410)

    def test_permission_and_question_responses_are_scoped(self):
        _,t=self.make_workspace()
        server=FakeRuntime.servers['http://alice']
        server['requests']['question']=[{"id":"q1","sessionID":t["session_id"],"questions":[]}]
        server['requests']['permission']=[{"id":"p1","sessionID":"other-session"}]
        ok=self.client.post(f'/api/threads/{t["id"]}/requests/question/q1',json={"answers":[["乙方"]]},headers=self.headers)
        self.assertEqual(ok.status_code,200)
        denied=self.client.post(f'/api/threads/{t["id"]}/requests/permission/p1',json={"reply":"once"},headers=self.headers)
        self.assertEqual(denied.status_code,404)

    def test_invalid_uploads(self):
        r=self.client.post('/api/workspaces',content=b'x',headers={**self.headers,'X-Filename':'test.exe'})
        self.assertEqual(r.status_code,415)
        r=self.client.post('/api/workspaces',content=b'x'*(20*1024*1024+1),headers={**self.headers,'X-Filename':'test.txt'})
        self.assertEqual(r.status_code,413)

    def test_failed_export_does_not_register_a_saved_artifact(self):
        w,t=self.make_workspace();self.read_all(w,t);body,_=self.report(w)
        with patch('contract_web.artifact_service._md_to_docx_bytes',side_effect=OSError('disk unavailable')):
            with self.assertRaises(OSError):self.publish(t,body)
        self.assertEqual(self.store.all('SELECT * FROM artifacts'),[])
        self.assertEqual(list((self.store.user_root(self.uid)/'published').iterdir()),[])

    def test_long_contract_must_be_read_across_all_chunks(self):
        w,t=self.make_workspace('\n'.join(f'第{i}条：金额{i}元' for i in range(2400)))
        body,m=self.report(w);self.read_all(w,t)
        tool=self.native(t)['messages'][0]['parts'][0]
        full=tool['state']['output'];tool['state']['output']='\n'.join(full.splitlines()[:2000])
        self.assertEqual(self.publish(t,body).status_code,422)
        remainder=copy.deepcopy(tool);remainder['id']='read-rest';remainder['state']['output']='\n'.join(full.splitlines()[2000:])
        self.native(t)['messages'][0]['parts'].append(remainder)
        self.assertEqual(self.publish(t,body).status_code,200)

    def test_native_read_rules_use_the_reported_worktree(self):
        w,t=self.make_workspace()
        reads=[r for r in self.native(t)['permission'] if r['permission']=='read' and r['action']=='allow']
        self.assertIn(f'sources/{w["document_id"]}/*',[r['pattern'] for r in reads])
        self.assertIn(f'work/{t["id"]}/*',[r['pattern'] for r in reads])

    def test_two_accounts_can_have_independent_running_sessions(self):
        wa,ta=self.make_workspace();self.assertEqual(self.send(ta).status_code,202)
        self.login('bob');wb,tb=self.make_workspace('乙方为其他公司。\n完全不同的合同，金额为 999999 元。')
        self.assertEqual(self.send(tb).status_code,202)
        self.assertEqual(self.client.get(f'/api/threads/{tb["id"]}').json()['status']['type'],'busy')
        self.assertEqual(self.client.get(f'/api/threads/{ta["id"]}').status_code,404)
        self.login('alice');self.assertEqual(self.client.get(f'/api/threads/{ta["id"]}').json()['status']['type'],'busy')

    def test_document_mapping_preserves_docx_table_order(self):
        doc=Document();doc.add_paragraph('第一条 总金额 128000 元');table=doc.add_table(rows=2,cols=2)
        table.cell(0,0).text='项目';table.cell(0,1).text='期限';table.cell(1,0).text='付款';table.cell(1,1).text='30天'
        doc.add_paragraph('第二条 验收');p=self.root/'example.docx';doc.save(p)
        mapped=prepare(p,'123456abcdef');text=(self.root/'contract.md').read_text()
        self.assertLess(text.index('128000'),text.index('付款'));self.assertLess(text.index('付款'),text.index('第二条'))
        self.assertTrue(any(s.get('skip') for s in mapped['segments']))

    def test_native_numbered_read_of_docx_soft_line_breaks(self):
        doc=Document();doc.add_paragraph('第一条 付款\n验收后 30 日内付款。')
        path=self.root/'soft-break.docx';doc.save(path);mapped=prepare(path,'123456abcdef')
        text=(self.root/'contract.md').read_text();output='\n'.join(f'{i+1}: {line}' for i,line in enumerate(text.splitlines()))
        message={'parts':[{'type':'tool','tool':'read','state':{'status':'completed','input':{'filePath':'/source'},'output':output}}]}
        self.assertEqual(read_blocks([message],{'123456abcdef':mapped},{'/source':'123456abcdef'})['123456abcdef'],{'B0'})

    def test_text_pdf_upload_mapping_page_render_and_ownership(self):
        pdf=(Path(__file__).parent/'fixtures/text-contract.pdf').read_bytes()
        w=self.client.post('/api/workspaces',content=pdf,headers={**self.headers,'X-Filename':'contract.pdf'}).json()
        t=self.client.post(f'/api/workspaces/{w["id"]}/threads',json={},headers=self.headers).json()
        path=f'/api/documents/{w["document_id"]}'
        mapped=self.client.get(path+'?thread_id='+t['id']).json()
        self.assertEqual(mapped['kind'],'pdf');self.assertEqual(mapped['segments'][0]['page'],1)
        self.assertIn('bbox',mapped['segments'][0])
        page=path+'/pages/1?thread_id='+t['id'];response=self.client.get(page)
        self.assertEqual(response.status_code,200);self.assertTrue(response.content.startswith(b'\x89PNG'))
        self.assertEqual(self.client.get(path+'/pages/0?thread_id='+t['id']).status_code,404)
        self.assertEqual(self.client.get(path+'/pages/2?thread_id='+t['id']).status_code,404)
        self.login('bob');self.assertEqual(self.client.get(page).status_code,404)


    def test_attachment_delete_revokes_access_and_preserves_other_materials(self):
        w,t=self.make_workspace()
        t2=self.client.post(f'/api/workspaces/{w["id"]}/threads',headers=self.headers,json={}).json()
        a=self.client.post(f'/api/threads/{t["id"]}/attachments',content=b'Private attachment',headers={**self.headers,"X-Filename":"extra.txt"}).json()
        delete=f'/api/threads/{t["id"]}/attachments/{a["id"]}'
        self.assertEqual(self.client.delete(f'/api/threads/{t2["id"]}/attachments/{a["id"]}',headers=self.headers).status_code,404)
        self.login('bob');self.assertEqual(self.client.delete(delete,headers=self.headers).status_code,404);self.login('alice')
        self.native(t)['status']={'type':'busy'}
        self.assertEqual(self.client.delete(delete,headers=self.headers).status_code,409)
        self.native(t)['status']={'type':'idle'}
        with patch.object(FakeRuntime,'permissions',side_effect=ValueError('offline')):
            self.assertEqual(self.client.delete(delete,headers=self.headers).status_code,422)
        self.assertTrue((self.store.user_root(self.uid)/'sources'/a['id']).is_dir())
        self.assertEqual(self.client.delete(delete,headers=self.headers).status_code,200)
        self.assertFalse((self.store.user_root(self.uid)/'sources'/a['id']).exists())
        self.assertFalse(any(a['id'] in rule['pattern'] for rule in self.native(t)['permission']))
        self.assertEqual(self.client.get(f'/api/documents/{a["id"]}/file?thread_id={t["id"]}').status_code,404)
        context=json.loads((self.store.user_root(self.uid)/'threads'/t['id']/'context.json').read_text())
        self.assertEqual([d['id'] for d in context['documents']],[w['document_id']])
        self.assertEqual(self.client.delete(f'/api/threads/{t["id"]}/attachments/{w["document_id"]}',headers=self.headers).status_code,422)

    def test_selected_quotes_validate_scope_version_and_real_text(self):
        w,t=self.make_workspace();_,m=self.report(w)
        q={'document_id':w['document_id'],'source_hash':m['source_hash'],'block_ids':['B2'],'text':'总金额 128000 元'}
        for bad in [{**q,'document_id':'other'},{**q,'source_hash':'stale'},{**q,'text':'总金额 999999 元'},{**q,'block_ids':['B0','B2']},{**q,'block_ids':[]}]:
            self.assertEqual(self.send(t,quotes=[bad]).status_code,422)
        self.assertEqual(self.send(t,'解释这一段',quotes=[q]).status_code,202)
        native_text=self.native(t)['messages'][-1]['parts'][0]['text']
        self.assertIn(q['text'],native_text);self.assertIn(m['segments'][2]['citation'],native_text)

    def test_extra_artifact_formats_have_exact_downloads_isolated_preview_and_ownership(self):
        w,t=self.make_workspace();body,_=self.report(w)
        contents={'html':'<!doctype html><h1>合同看板</h1><script>document.body.dataset.ready="yes"</script>',
                  'txt':'付款计划\n第一期 30%', 'json':'{"amount":128000}', 'csv':'项目,金额\n"服务,一期",128000',
                  'svg':'<svg xmlns="http://www.w3.org/2000/svg"><text x="10" y="20">合同</text></svg>'}
        for fmt,content in contents.items():
            saved=self.publish(t,{**body,'kind':'document','format':fmt,'content':content,'id':'untrusted-id','thread_id':'wrong-thread'})
            self.assertEqual(saved.status_code,200,saved.text);aid=saved.json()['artifact_id']
            detail=self.client.get(f'/api/artifacts/{aid}').json()
            self.assertEqual(detail['formats'],[fmt]);self.assertEqual(detail['content'],content)
            self.assertEqual(detail['id'],aid);self.assertEqual(detail['thread_id'],t['id'])
            download=self.client.get(f'/api/artifacts/{aid}/file?format={fmt}')
            self.assertEqual(download.content,content.encode());self.assertIn('attachment',download.headers['content-disposition'])
            self.assertEqual(self.client.get(f'/api/artifacts/{aid}/file?format=docx').status_code,422)
            preview=self.client.get(f'/api/artifacts/{aid}/preview')
            self.assertEqual(preview.status_code,200 if fmt in {'html','svg'} else 422)
            if fmt in {'html','svg'}:
                self.assertEqual(preview.content,download.content)
                self.assertIn("sandbox allow-scripts",preview.headers['content-security-policy'])
                self.assertIn("connect-src 'none'",preview.headers['content-security-policy'])
                self.assertNotIn('allow-same-origin',preview.headers['content-security-policy'])
            self.login('bob');self.assertEqual(self.client.get(f'/api/artifacts/{aid}/preview').status_code,404);self.login('alice')
        for fmt,content in [('exe','no'),('json','broken'),('svg','<div>wrong</div>')]:
            self.assertEqual(self.publish(t,{**body,'kind':'document','format':fmt,'content':content}).status_code,422)
        # HTML does not bypass the complete review/summary evidence checks.
        self.assertEqual(self.publish(t,{**body,'format':'html'}).status_code,422)
        self.read_all(w,t)
        self.assertEqual(self.publish(t,{**body,'format':'html'}).status_code,200)

    def test_native_paths_and_raw_tool_data_are_not_sent_in_snapshot_or_events(self):
        from contract_web.presentation import PublicView
        w,t=self.make_workspace();doc=w['document_id']
        path=f'/sources/{doc}/contract.md'
        part={'id':'p','messageID':'m','sessionID':t['session_id'],'type':'tool','tool':'read','state':{
            'status':'completed','input':{'filePath':path,'offset':1533,'limit':180},
            'output':'PRIVATE RAW OUTPUT /private/var/folders/secret','metadata':{'filePath':path}}}
        self.native(t)['messages']=[{'info':{'id':'m','role':'assistant','path':{'cwd':'/private/var/secret'}},'parts':[part]}]
        visible=self.client.get(f'/api/threads/{t["id"]}').json();encoded=json.dumps(visible)
        self.assertNotIn('/private',encoded);self.assertNotIn('/sources',encoded);self.assertNotIn('PRIVATE RAW',encoded)
        self.assertIn('contract.txt',encoded);self.assertIn('1533',encoded);self.assertIn('1712',encoded)
        self.assertIn('filePath',self.native(t)['messages'][0]['parts'][0]['state']['input'])
        view=PublicView({path:'contract.txt'})
        event=visible_event({'type':'message.part.updated','properties':{'part':part}},t['session_id'],set(),view)
        self.assertNotIn('filePath',json.dumps(event));self.assertNotIn('PRIVATE RAW',json.dumps(event))
        known=set();text={'id':'text','messageID':'m','sessionID':t['session_id'],'type':'text','text':'读取 /private/var/'}
        visible_event({'type':'message.part.updated','properties':{'part':text}},t['session_id'],known,view)
        event=visible_event({'type':'message.part.delta','properties':{'sessionID':t['session_id'],'messageID':'m','partID':'text','field':'text','delta':'folders/user/contract.md'}},t['session_id'],known,view)
        self.assertNotIn('/private',json.dumps(event));self.assertNotIn('folders/user',json.dumps(event))

    def test_html_preview_citations_use_current_workspace_and_keep_download_bytes(self):
        w,t=self.make_workspace();body,_=self.report(w)
        content='<h1>原文引用【D'+w['document_id']+':B0】</h1>'
        saved=self.publish(t,{**body,'kind':'document','format':'html','content':content})
        self.assertEqual(saved.status_code,200,saved.text);aid=saved.json()['artifact_id']
        response=self.client.get('/api/artifacts/'+aid+'/preview',params={'thread_id':t['id']})
        self.assertEqual(response.status_code,200);self.assertIn('data-workbench-reference="0"',response.text)
        self.assertNotIn('allow-same-origin',response.headers['content-security-policy'])
        self.assertEqual(self.client.get('/api/artifacts/'+aid+'/file?format=html').text,content)
        other,ot=self.make_workspace()
        self.assertEqual(self.client.get('/api/artifacts/'+aid+'/preview',params={'thread_id':ot['id']}).status_code,404)

    def test_snapshot_includes_preview_only_for_uniquely_resolved_references(self):
        w,t=self.make_workspace();did=w['document_id']
        self.native(t)['messages']=[{'info':{'id':'a','role':'assistant','sessionID':t['session_id'],'time':{'completed':1}},'parts':[{'id':'p','type':'text','text':'合同说明【B0】'}]}]
        snap=self.client.get('/api/threads/'+t['id']).json()
        self.assertTrue(snap['documents'][0]['locations']['B0']['preview'])
        self.assertFalse(snap['full_execution_available'])
        self.client.post('/api/threads/'+t['id']+'/attachments',content=b'another source',headers={**self.headers,'X-Filename':'other.txt'}).raise_for_status()
        snap=self.client.get('/api/threads/'+t['id']).json()
        self.assertTrue(all('preview' not in d['locations']['B0'] for d in snap['documents']))
        self.native(t)['messages'][0]['parts'][0]['text']='合同说明【D'+did+':B0】'
        snap=self.client.get('/api/threads/'+t['id']).json()
        self.assertTrue(next(d for d in snap['documents'] if d['id']==did)['locations']['B0']['preview'])

    def test_permission_modes_persist_and_keep_scope(self):
        w,t=self.make_workspace();url=f'/api/threads/{t["id"]}/permission-mode'
        def change(mode):return self.client.put(url,json={'mode':mode},headers=self.headers)
        def edit_action(pattern):return next(r['action'] for r in self.native(t)['permission'] if r['permission']=='edit' and r['pattern']==pattern)
        self.assertEqual(change('cautious').status_code,200)
        self.assertEqual(edit_action(f'work/{t["id"]}/*'),'ask')
        self.assertEqual(edit_action(f'work/{t["id"]}/context.json'),'deny')
        self.assertEqual(self.client.get(f'/api/threads/{t["id"]}').json()['permission_mode'],'cautious')
        self.assertEqual(change('bypass').status_code,422)
        self.assertEqual(change('full').status_code,422)
        self.native(t)['status']={'type':'busy'};self.assertEqual(change('auto').status_code,409)
        self.native(t)['status']={'type':'idle'}
        server=FakeRuntime.servers['http://alice'];server['requests']['permission']=[{'id':'p','sessionID':t['session_id']}]
        self.assertEqual(change('auto').status_code,409)
        response=self.client.post(f'/api/threads/{t["id"]}/requests/permission/p',json={'reply':'always'},headers=self.headers)
        self.assertEqual(response.status_code,200)
        self.assertEqual(server['calls'][-1][3],{'reply':'always'})
        server['requests']['permission']=[]
        a=self.client.post(f'/api/threads/{t["id"]}/attachments',content=b'extra',headers={**self.headers,'X-Filename':'extra.txt'})
        self.assertEqual(a.status_code,200)
        self.assertEqual(edit_action(f'work/{t["id"]}/*'),'ask')
        self.assertEqual(change('auto').status_code,200)
        self.assertEqual(edit_action(f'work/{t["id"]}/*'),'allow')
        self.assertEqual(edit_action(f'work/{t["id"]}/.publish-token'),'deny')
        self.login('bob');self.assertEqual(change('auto').status_code,404)

    def review_fixture(self):
        w,t=self.make_workspace();self.read_all(w,t);body,m=self.report(w)
        cites=''.join(s['citation'] for s in m['segments'][2:4])
        body.update(kind='review',content='TEST-1 '+cites,findings=[{'risk_id':'TEST-1','verdict':'命中','reason':'审查依据'+cites,'evidence':[
            {'document_id':w['document_id'],'source_hash':m['source_hash'],'block_id':s['id'],'quote':s['text']} for s in m['segments'][2:4]]}])
        saved=self.publish(t,body);self.assertEqual(saved.status_code,200,saved.text)
        return w,t,body,saved.json()['artifact_id']

    def test_review_normalized_files_and_legacy_projection(self):
        w,t,body,aid=self.review_fixture()
        data=self.client.get(f'/api/artifacts/{aid}').json()
        self.assertIn(':B2-B3',data['content']);self.assertEqual(data['findings'][0]['risk_name'],'付款')
        self.assertEqual(len(data['findings'][0]['evidence']),2)
        self.assertEqual(len(data['findings'][0]['evidence_groups']),1)
        download=self.client.get(f'/api/artifacts/{aid}/file?format=md').text
        self.assertEqual(download,data['content'])
        doc=Document(io.BytesIO(self.client.get(f'/api/artifacts/{aid}/file?format=docx').content))
        text='\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('第 3–4 段',text);self.assertNotIn(w['document_id'],text)
        dest=self.store.user_root(self.uid)/'published'/aid
        (dest/'report.json').write_text(json.dumps(body));before=(dest/'report.json').read_bytes()
        self.assertIn(':B2-B3',self.client.get(f'/api/artifacts/{aid}').json()['content'])
        self.assertEqual(before,(dest/'report.json').read_bytes())

    def test_risk_feedback_versions_history_ownership_and_independent_storage(self):
        w,t,report,aid=self.review_fixture();url=f'/api/artifacts/{aid}/risks/TEST-1/feedback'
        body={'source_hash':report['source_hash'],'decision':'接受风险','note':'金额可控，由项目经理确认','revision':0}
        save=lambda b:self.client.put(url,json=b,headers=self.headers)
        self.assertEqual(save({**body,'note':''}).status_code,422)
        self.assertEqual(save({**body,'source_hash':'old'}).status_code,409)
        first=save(body);self.assertEqual(first.status_code,200,first.text)
        self.assertEqual(first.json()['revision'],1)
        self.assertEqual(save(body).status_code,409)
        second=save({**body,'revision':1,'decision':'判断有误','note':'请核对付款前置条件'})
        self.assertEqual(second.status_code,200);self.assertEqual(second.json()['history'][0]['decision'],'接受风险')
        data=self.client.get(f'/api/artifacts/{aid}').json();self.assertEqual(data['findings'][0]['verdict'],'命中')
        with TestClient(create_app(self.root,FakeRuntime,self.library)) as restarted:
            restarted.cookies.update(self.client.cookies)
            self.assertEqual(restarted.get(f'/api/artifacts/{aid}').json()['feedback']['TEST-1']['revision'],2)
        self.client.post(f'/api/threads/{t["id"]}/messages',json={'text':'按意见处理'},headers=self.headers)
        context=json.loads((self.store.user_root(self.uid)/'threads'/t['id']/'context.json').read_text())
        self.assertNotIn('risk_feedback',context)
        t2=self.client.post(f'/api/workspaces/{w["id"]}/threads',json={},headers=self.headers).json()
        self.assertNotIn('risk_feedback',json.loads((self.store.user_root(self.uid)/'threads'/t2['id']/'context.json').read_text()))
        self.login('bob');self.assertEqual(save({**body,'revision':2}).status_code,404)
        self.login('alice');self.store.execute('UPDATE documents SET source_hash=? WHERE id=?',('changed',w['document_id']))
        self.assertEqual(save({**body,'revision':2}).status_code,409)

if __name__ == '__main__':
    unittest.main()
