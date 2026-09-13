"""Document export extracted from the original workbench; no model calls."""
import io
import re

def _strip_markdown_inline(value: str) -> str:
    text = re.sub(r"!\[([^]]*)\]\([^)]*\)", r"\1", str(value or ""))
    text = re.sub(r"\[([^]]+)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    return re.sub(r"[*_~`]", "", text).strip()

def _md_to_docx_bytes(md: str) -> bytes:
    """Markdown -> Word with headings, lists, quotes and pipe tables."""
    import docx

    doc = docx.Document()
    table_buf: list[list[str]] = []

    def flush_table():
        rows = [row for row in table_buf if not all(
            re.fullmatch(r"[-:\s]*", cell) for cell in row)]
        table_buf.clear()
        if not rows or not rows[0]:
            return
        cols = max(len(row) for row in rows)
        table = doc.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for i, row in enumerate(rows):
            for j, cell in enumerate(row):
                table.cell(i, j).text = _strip_markdown_inline(cell)

    for raw in str(md or "").splitlines():
        line = raw.rstrip()
        if re.match(r"^\|.*\|$", line.strip()):
            table_buf.append([cell.strip() for cell in line.strip()[1:-1].split("|")])
            continue
        flush_table()
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        numbered = re.match(r"^\s*\d+[.)、]\s+(.+)$", line)
        if heading:
            doc.add_heading(_strip_markdown_inline(heading.group(2)), len(heading.group(1)))
        elif bullet:
            doc.add_paragraph(_strip_markdown_inline(bullet.group(1)), style="List Bullet")
        elif numbered:
            doc.add_paragraph(_strip_markdown_inline(numbered.group(1)), style="List Number")
        elif line.startswith("> "):
            doc.add_paragraph(_strip_markdown_inline(line[2:]), style="Intense Quote")
        elif line.strip() and line.strip() != "---":
            doc.add_paragraph(_strip_markdown_inline(line))
    flush_table()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
