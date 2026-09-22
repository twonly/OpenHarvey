from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from PIL import Image
import argparse
parser=argparse.ArgumentParser();parser.add_argument('--output',default='.cache/redline-corpus');args=parser.parse_args()
root=Path(args.output);root.mkdir(parents=True,exist_ok=True); Image.new('RGB',(160,50),'navy').save(root/'logo.png')
names=['chinese','english','mixed_runs','numbered','table','merged_table','header_footer','sections','image','long_60_pages','comments','bilingual']
for name in names:
 d=Document();d.add_heading('合同兼容性样本 '+name,0)
 p=d.add_paragraph();p.add_run('买方付款期限为 ');p.add_run('30').bold=name=='mixed_runs';p.add_run(' 天。 Payment term: 30 days.')
 if name=='english': d.add_paragraph('The supplier warrants the goods for twelve months.')
 if name=='numbered':
  for i in range(6): d.add_paragraph('付款义务与通知要求 '+str(i),style='List Number' if i%2==0 else 'List Number 2')
 if name in ['table','merged_table']:
  t=d.add_table(rows=130 if name=='table' else 5, cols=3);t.style='Table Grid'
  for i,r in enumerate(t.rows):
   for j,c in enumerate(r.cells):c.text=f'交付项 {i}-{j}'
  if name=='merged_table':t.cell(0,0).merge(t.cell(0,1))
 if name=='header_footer':
  d.sections[0].header.paragraphs[0].text='测试合同 · 保密';d.sections[0].footer.paragraphs[0].text='合同编号 TEST-2026'
 if name=='sections':d.add_section(WD_SECTION.NEW_PAGE);d.add_paragraph('附件：服务范围')
 if name=='image': d.add_picture(str(root/'logo.png'),width=Inches(1.5))
 if name=='long_60_pages':
  for i in range(60):d.add_page_break();d.add_heading(f'第 {i+1} 项 服务约定',1);d.add_paragraph('本段用于分页兼容性测试。'*40)
 if name=='comments':d.add_comment(p.runs,text='请确认付款日期',author='原审阅人',initials='原')
 if name=='bilingual':d.add_paragraph('适用法律与争议解决 / Governing Law and Dispute Resolution')
 d.add_paragraph('KEEP_UNCHANGED_验收与其他条款。')
 d.save(root/(name+'.docx'))
print('Created 12 synthetic DOCX fixtures in',root)
